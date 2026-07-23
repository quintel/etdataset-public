"""
UNFCCC Excel CRT submission data source.

Parses UNFCCC Common Reporting Tables (CRT) Excel files for countries that
don't report via API (e.g., Serbia, Singapore, Netherlands 2023).
"""

from __future__ import annotations
from typing import Iterator
from pathlib import Path
import pandas as pd
import openpyxl
import re

from .base import DataSource
from ..config import MAX_CATEGORY_DEPTH, GWP_AR5_CH4, GWP_AR5_N2O

# Matches CRT category labels: "1. B. 1. b. Fuel transformation" → code="1.B.1.b", name="Fuel transformation"
_CRT_CODE_RE = re.compile(
    r"^([\d]+(?:\.\s*(?:[A-Z]|[\d]+|[a-z]{1,2}|i{1,3}|iv|v|vi{1,3}|ix|x)\b)*)\.?\s*(.*)"
)
_CRT_FOOTNOTE_RE = re.compile(r"\s*\(\d+\)\s*$")


class UFCCCExcelSource(DataSource):
    """Parse UNFCCC CRT Excel submissions."""

    def fetch(
        self,
        dataset: str,
        years: list[int],
        excel_path: Path,
        max_depth: int = MAX_CATEGORY_DEPTH,
    ) -> Iterator[dict]:
        """
        Parse Excel file and yield emission records.

        Args:
            dataset: Dataset identifier
            years: Years to fetch (should have one entry matching file year)
            excel_path: Path to CRT Excel file
            max_depth: Maximum CRT code depth

        Yields:
            Standard emission records
        """
        year = years[0] if years else None
        if not year:
            raise ValueError("Year must be provided for Excel source")

        df = self._parse_excel_file(excel_path, year, max_depth)

        for _, row in df.iterrows():
            yield self._standardize_row(
                code=row["category_code"],
                category=row["category"],
                year=row["year"],
                co2=row["co2"],
                other_ghg=row["other_ghg_co2eq"],
            )

    def _parse_excel_file(
        self,
        path: Path,
        year: int,
        max_depth: int,
    ) -> pd.DataFrame:
        """Parse all relevant sheets from Excel file."""
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        all_rows: list[dict] = []

        self._parse_summary2(wb, year, max_depth, all_rows)
        self._parse_energy_sheets(wb, year, max_depth, all_rows)
        self._parse_sector_sheets(wb, year, max_depth, all_rows)

        wb.close()

        return self._deduplicate_rows(all_rows)

    def _parse_summary2(
        self,
        wb,
        year: int,
        max_depth: int,
        all_rows: list[dict],
    ) -> None:
        """Parse Summary2 sheet (aggregate totals)."""
        self._parse_summary2_ws(wb["Summary2"], year, max_depth, all_rows)

    @classmethod
    def _parse_summary2_ws(
        cls,
        ws,
        year: int,
        max_depth: int,
        all_rows: list[dict],
    ) -> None:
        """
        Parse a Summary2 worksheet into emission rows.

        Shared by the fetch pipeline and Summary2 validation so the two never
        diverge in how the sheet is read.
        """
        # max_col=10 covers CO2 (col 2) plus the seven other-GHG columns 3-9
        # (CH4, N2O, HFCs, PFCs, Unspecified mix, SF6, NF3); the "Total" column
        # sits beyond this range and is intentionally excluded.
        for i, row in enumerate(ws.iter_rows(min_col=1, max_col=10, values_only=True)):
            if i < 9:
                continue

            label = str(row[1]).strip() if row[1] is not None else ""

            if label.startswith("Indirect CO2"):
                all_rows.append(
                    {
                        "category_code": "ind_CO2",
                        "category": "Indirect CO2",
                        "year": year,
                        "co2": cls._num(row[2]),
                        "other_ghg_co2eq": float("nan"),
                        "_source_priority": 0,
                    }
                )
                continue

            code, name = cls._extract_crt_code(label)
            if not code or code.count(".") > max_depth:
                continue

            co2 = cls._num(row[2])
            other_vals = [cls._num(v) for v in row[3:10]]
            other = cls._sum_numeric(other_vals)

            all_rows.append(
                {
                    "category_code": code,
                    "category": name,
                    "year": year,
                    "co2": co2,
                    "other_ghg_co2eq": other,
                    "_source_priority": 0,
                }
            )

    # Top-level CRT sectors summed for Summary2 validation. LULUCF (sector 4)
    # and sector 6 are excluded, matching "total GHG emissions without LULUCF".
    SUMMARY2_TOTAL_SECTORS = ("1", "2", "3", "5")

    @classmethod
    def summary2_totals(
        cls,
        excel_path: Path,
        year: int,
        sectors: tuple[str, ...] = SUMMARY2_TOTAL_SECTORS,
    ) -> tuple[float, float] | None:
        """
        Sum Summary2 CO2 and other-GHG over the given top-level CRT sectors.

        Reuses the same Summary2 parsing as the fetch pipeline so validation and
        ingestion never diverge. Non-numeric cells count as 0.

        Args:
            excel_path: Path to the UNFCCC CRT Excel file
            year: Year to record on the parsed rows
            sectors: Top-level CRT sector codes to sum (default: 1, 2, 3, 5)

        Returns:
            (co2_total, other_ghg_total) in kton CO2-eq, or None when the file
            cannot be opened or has no Summary2 sheet.
        """
        try:
            wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
        except Exception:
            return None

        if "Summary2" not in wb.sheetnames:
            wb.close()
            return None

        all_rows: list[dict] = []
        cls._parse_summary2_ws(wb["Summary2"], year, MAX_CATEGORY_DEPTH, all_rows)
        wb.close()

        sector_set = set(sectors)
        co2_total = 0.0
        other_total = 0.0
        for r in all_rows:
            if r["category_code"] not in sector_set:
                continue
            co2 = r["co2"]
            other = r["other_ghg_co2eq"]
            co2_total += 0.0 if pd.isna(co2) else co2
            other_total += 0.0 if pd.isna(other) else other

        return co2_total, other_total

    def _parse_energy_sheets(
        self,
        wb,
        year: int,
        max_depth: int,
        all_rows: list[dict],
    ) -> None:
        """Parse Energy sector sheets (Table1.A, 1.B, 1.C)."""
        # Table1.A(a)s1-4: combustion by sub-sector
        for sheet in [
            "Table1.A(a)s1",
            "Table1.A(a)s2",
            "Table1.A(a)s3",
            "Table1.A(a)s4",
        ]:
            if sheet in wb.sheetnames:
                self._parse_sheet(
                    wb[sheet],
                    year,
                    max_depth,
                    all_rows,
                    source_priority=1,
                    co2_idx=7,
                    ch4_kt_idx=8,
                    n2o_kt_idx=9,
                    max_col=11,
                )

        # Table1.B.1: fugitive emissions from solid fuels
        if "Table1.B.1" in wb.sheetnames:
            self._parse_sheet(
                wb["Table1.B.1"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=6,
                ch4_kt_idx=5,
                max_col=10,
            )

        # Table1.B.2: fugitive emissions from oil & gas
        if "Table1.B.2" in wb.sheetnames:
            self._parse_sheet(
                wb["Table1.B.2"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=8,
                ch4_kt_idx=9,
                n2o_kt_idx=10,
                max_col=12,
            )

        # Table1.C: CO2 transport & storage
        if "Table1.C" in wb.sheetnames:
            self._parse_sheet(
                wb["Table1.C"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=4,
                max_col=6,
            )

    def _parse_sector_sheets(
        self,
        wb,
        year: int,
        max_depth: int,
        all_rows: list[dict],
    ) -> None:
        """Parse other sector sheets (Table2-5)."""
        # Table2(I): Industrial Processes
        if "Table2(I)" in wb.sheetnames:
            self._parse_sheet(
                wb["Table2(I)"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=2,
                total_co2eq_idx=14,
                max_col=15,
            )

        # Table3: Agriculture
        if "Table3" in wb.sheetnames:
            self._parse_sheet(
                wb["Table3"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=2,
                total_co2eq_idx=9,
                max_col=11,
            )

        # Table4: LULUCF
        if "Table4" in wb.sheetnames:
            self._parse_sheet(
                wb["Table4"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=2,
                total_co2eq_idx=8,
                max_col=10,
            )

        # Table5: Waste
        if "Table5" in wb.sheetnames:
            self._parse_sheet(
                wb["Table5"],
                year,
                max_depth,
                all_rows,
                source_priority=1,
                co2_idx=2,
                total_co2eq_idx=9,
                max_col=11,
            )

    def _parse_sheet(
        self,
        ws,
        year: int,
        max_depth: int,
        all_rows: list[dict],
        source_priority: int,
        co2_idx: int,
        other_sum_idx: list[int] | None = None,
        total_co2eq_idx: int | None = None,
        ch4_kt_idx: int | None = None,
        n2o_kt_idx: int | None = None,
        skip_rows: int = 9,
        max_col: int = 11,
    ) -> None:
        """Parse a single sheet and append rows."""
        for i, row in enumerate(
            ws.iter_rows(min_col=1, max_col=max_col, values_only=True)
        ):
            if i < skip_rows:
                continue

            label = str(row[1]).strip() if row[1] is not None else ""
            code, name = self._extract_crt_code(label)

            if not code:
                continue

            top_level = code.split(".")[0]
            if top_level.isdigit() and int(top_level) > 6:
                continue

            if code.count(".") > max_depth:
                continue

            co2 = self._num(row[co2_idx])
            other = self._calculate_other_ghg(
                row, other_sum_idx, total_co2eq_idx, ch4_kt_idx, n2o_kt_idx, co2
            )

            all_rows.append(
                {
                    "category_code": code,
                    "category": name,
                    "year": year,
                    "co2": co2,
                    "other_ghg_co2eq": other,
                    "_source_priority": source_priority,
                }
            )

    def _calculate_other_ghg(
        self,
        row,
        other_sum_idx: list[int] | None,
        total_co2eq_idx: int | None,
        ch4_kt_idx: int | None,
        n2o_kt_idx: int | None,
        co2: float,
    ) -> float:
        """Calculate other_ghg_co2eq from different source columns."""
        # Priority 1: Direct CO2eq columns
        if other_sum_idx is not None:
            other_vals = [self._num(row[c]) for c in other_sum_idx]
            return self._sum_numeric(other_vals)

        # Priority 2: Total GHG CO2eq minus CO2
        if total_co2eq_idx is not None:
            total = self._num(row[total_co2eq_idx])
            if pd.isna(total) and pd.isna(co2):
                return float("nan")
            total_val = 0.0 if pd.isna(total) else total
            co2_val = 0.0 if pd.isna(co2) else co2
            return total_val - co2_val

        # Priority 3: Raw CH4/N2O kt converted to CO2eq
        if ch4_kt_idx is not None or n2o_kt_idx is not None:
            ch4_kt = (
                self._num(row[ch4_kt_idx]) if ch4_kt_idx is not None else float("nan")
            )
            n2o_kt = (
                self._num(row[n2o_kt_idx]) if n2o_kt_idx is not None else float("nan")
            )

            if pd.isna(ch4_kt) and (n2o_kt_idx is None or pd.isna(n2o_kt)):
                return float("nan")

            ch4_co2eq = ch4_kt * GWP_AR5_CH4 if pd.notna(ch4_kt) else 0.0
            n2o_co2eq = n2o_kt * GWP_AR5_N2O if pd.notna(n2o_kt) else 0.0
            return ch4_co2eq + n2o_co2eq

        return float("nan")

    def _deduplicate_rows(
        self,
        all_rows: list[dict],
    ) -> pd.DataFrame:
        """Deduplicate rows, preferring detailed sheets over Summary2."""
        result = pd.DataFrame(all_rows)

        result["_completeness"] = result["co2"].notna().astype(int) + result[
            "other_ghg_co2eq"
        ].notna().astype(int)
        result["_score"] = result["_source_priority"] * 10 + result["_completeness"]

        result = (
            result.sort_values(by="_score", ascending=False)
            .drop_duplicates(subset=["category_code", "year"], keep="first")
            .drop(columns=["_source_priority", "_completeness", "_score"])
            .reset_index(drop=True)
        )

        return result[["category_code", "category", "year", "co2", "other_ghg_co2eq"]]

    @staticmethod
    def _num(v) -> float:
        """Convert value to numeric, reporting flags become NaN."""
        return v if isinstance(v, (int, float)) else float("nan")

    @staticmethod
    def _sum_numeric(values: list[float]) -> float:
        """Sum numeric values, return NaN if all missing."""
        numeric_vals = [v for v in values if pd.notna(v)]
        return sum(numeric_vals) if numeric_vals else float("nan")

    @staticmethod
    def _extract_crt_code(label: str) -> tuple[str, str]:
        """Extract CRT code and category name from label."""
        m = _CRT_CODE_RE.match(label)
        if not m:
            return "", ""

        code = m.group(1).replace(" ", "").rstrip(".")
        name = _CRT_FOOTNOTE_RE.sub("", m.group(2)).strip()
        return code, name
